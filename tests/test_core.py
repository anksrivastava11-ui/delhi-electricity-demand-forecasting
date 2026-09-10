import numpy as np
import pandas as pd
from openpyxl import load_workbook

from electricity_forecasting.features import make_supervised
from electricity_forecasting.intervals import ConformalInterval
from electricity_forecasting.metrics import forecast_metrics
from electricity_forecasting.model_comparison import export_model_comparison
from electricity_forecasting.sensitivity import select_social_feature_columns
from electricity_forecasting.split import last_year_train_test_split
from electricity_forecasting.data import _find_year_file, yearly_data_quality_report
from electricity_forecasting.pipeline import _save_social_analysis_unavailable
from electricity_forecasting.config import ForecastConfig


def test_lag_features_do_not_use_current_target():
    data = pd.DataFrame({"Date": pd.date_range("2023-01-01", periods=10), "demand": range(10)})
    featured = make_supervised(data, "demand", (1,), (3,))
    assert featured.loc[4, "lag_1"] == 3
    assert featured.loc[4, "rolling_mean_3"] == 2


def test_last_year_split_keeps_order():
    data = pd.DataFrame({"Date": pd.date_range("2022-01-01", "2024-01-01"), "demand": range(731)})
    train, test = last_year_train_test_split(data)
    assert train.Date.max() < test.Date.min()
    assert len(test) == 365


def test_conformal_interval_is_symmetric():
    interval = ConformalInterval(.9).fit(np.arange(20), np.arange(20) + 2)
    lower, upper = interval.predict(np.array([10.0]))
    assert lower[0] == 8 and upper[0] == 12


def test_metrics_include_required_error_measures():
    metrics = forecast_metrics([1, 2, 3], [1, 2, 4])
    assert {"MAE", "MSE", "RMSE", "R2", "MAPE"}.issubset(metrics)


def test_model_comparison_exports_separate_periods(tmp_path):
    metrics = pd.DataFrame([
        {"model": "better", "split": "train", "observations": 3, "MAE": 1, "MSE": 1, "RMSE": 1, "MAPE": 1, "R2": 0.9, "MASE": 1},
        {"model": "worse", "split": "train", "observations": 3, "MAE": 2, "MSE": 4, "RMSE": 2, "MAPE": 2, "R2": 0.1, "MASE": 2},
        {"model": "better", "split": "test", "observations": 2, "MAE": 3, "MSE": 9, "RMSE": 3, "MAPE": 3, "R2": 0.2, "MASE": 3},
        {"model": "worse", "split": "test", "observations": 2, "MAE": 4, "MSE": 16, "RMSE": 4, "MAPE": 4, "R2": 0.1, "MASE": 4},
    ])
    metrics_path = tmp_path / "metrics.csv"
    workbook_path = tmp_path / "comparison.xlsx"
    metrics.to_csv(metrics_path, index=False)

    outputs = export_model_comparison(metrics_path, tmp_path / "charts", workbook_path)

    assert outputs["train_scores"].iloc[0]["model"] == "better"
    assert outputs["test_scores"].iloc[0]["model"] == "better"
    assert len(list((tmp_path / "charts").glob("*.png"))) == 14
    assert set(load_workbook(workbook_path, read_only=True).sheetnames) == {
        "Methodology", "Train Metrics", "Train Scores", "Test Metrics", "Test Scores"
    }


def test_select_social_feature_columns_filters_social_indicators():
    features = [
        "lag_7",
        "Population in Rural Area",
        "Month",
        "Gross State Domestic Product (Current Prices)",
        "Max_Temperature",
    ]
    selected = select_social_feature_columns(features)
    assert selected == [
        "Population in Rural Area",
        "Gross State Domestic Product (Current Prices)",
    ]


def test_yearly_quality_report_exposes_missing_metadata():
    data = pd.DataFrame({
        "Date": pd.to_datetime(["2018-01-01", "2018-01-02", "2019-01-01"]),
        "demand": [1, 2, 3],
        "Summer": [np.nan, np.nan, 1],
    })
    report = yearly_data_quality_report(data)
    row = report.query("year == 2018 and column == 'Summer'").iloc[0]
    assert row.missing_values == 2
    assert row.missing_percent == 100


def test_workbook_selection_prefers_richer_duplicate_year_file(tmp_path):
    minimal = pd.DataFrame({"Date": [1], "Unrestricted demand": [1], "Max_Temperature": [2], "Min_Temperature": [1]})
    rich = minimal.assign(Summer=1, Winter=0)
    with pd.ExcelWriter(tmp_path / "2018 Manav.xlsx") as writer:
        minimal.to_excel(writer, sheet_name="Jan", index=False)
    with pd.ExcelWriter(tmp_path / "2018.xlsx") as writer:
        rich.to_excel(writer, sheet_name="Jan", index=False)
    selected = _find_year_file(tmp_path, 2018, required_columns=("Unrestricted demand", "Max_Temperature", "Min_Temperature"))
    assert selected.name == "2018.xlsx"


def test_unavailable_social_analysis_writes_transparent_artifacts(tmp_path):
    config = ForecastConfig(artifacts_dir=tmp_path)
    _save_social_analysis_unavailable("lstm", config, "LSTM has no social inputs.")
    report = pd.read_csv(tmp_path / "sensitivity_lstm_social.csv")
    assert report.loc[0, "status"] == "unavailable"
    assert (tmp_path / "feature_importance_lstm_social.png").exists()
    assert (tmp_path / "sensitivity_lstm_social.png").exists()
